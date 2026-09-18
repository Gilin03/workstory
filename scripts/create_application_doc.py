import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "deliverables" / "김태빈_지원서_묶음.docx"
APPROVED = json.loads((ROOT / "src" / "content" / "approved.json").read_text(encoding="utf-8"))


def set_cell_shading(cell, fill):
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_run_font(run, name="Malgun Gothic", size=10.5, color=None, bold=False):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def add_text(paragraph, text, **kwargs):
    run = paragraph.add_run(text)
    set_run_font(run, **kwargs)
    return run


def add_heading(doc, text, level=1):
    paragraph = doc.add_paragraph()
    paragraph.style = f"Heading {level}"
    paragraph.paragraph_format.space_before = Pt(16 if level == 1 else 10)
    paragraph.paragraph_format.space_after = Pt(6)
    add_text(paragraph, text, size=16 if level == 1 else 12.5, color=(225, 105, 68) if level == 1 else (23, 33, 39), bold=True)
    return paragraph


def add_body(doc, text, bold_prefix=None):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(7)
    paragraph.paragraph_format.line_spacing = 1.28
    if bold_prefix and text.startswith(bold_prefix):
        add_text(paragraph, bold_prefix, bold=True)
        add_text(paragraph, text[len(bold_prefix):])
    else:
        add_text(paragraph, text)
    return paragraph


def add_bullet(doc, text):
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(4)
    add_text(paragraph, text, size=10)
    return paragraph


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        set_cell_shading(cell, "0D1821")
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.paragraph_format.space_after = Pt(0)
        add_text(paragraph, header, size=8.5, color=(251, 250, 247), bold=True)
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            set_cell_margins(cells[index])
            cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            paragraph = cells[index].paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.15
            add_text(paragraph, value, size=8.5)
    if widths:
        for row in table.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = Cm(width)
    return table


def main():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(1.7)
    section.bottom_margin = Cm(1.7)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)

    styles = doc.styles
    styles["Normal"].font.name = "Malgun Gothic"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
    styles["Normal"].font.size = Pt(10.5)
    styles["Title"].font.name = "Malgun Gothic"
    styles["Title"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.paragraph_format.space_after = Pt(3)
    add_text(title, "김태빈", size=27, color=(13, 24, 33), bold=True)
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(16)
    add_text(subtitle, "지원 서류 묶음  |  웹 애플리케이션 · 정보보안", size=11, color=(102, 112, 123))

    contact = doc.add_table(rows=1, cols=3)
    contact.alignment = WD_TABLE_ALIGNMENT.LEFT
    contact.style = "Table Grid"
    contact_value = APPROVED["contact"]["email"]
    contact_note = "제출 전 공개용 전용 이메일로 교체" if APPROVED["contact"].get("isPlaceholder") else "공개하기로 정한 연락 수단"
    for cell, value in zip(contact.rows[0].cells, ["연락처", contact_value, contact_note]):
        set_cell_margins(cell, top=80, bottom=80)
        add_text(cell.paragraphs[0], value, size=9, color=(225, 105, 68) if cell == contact.rows[0].cells[0] else (23, 33, 39), bold=cell == contact.rows[0].cells[0])

    add_heading(doc, "1. 이력서", 1)
    add_body(doc, f"{APPROVED['tagline']}.")
    add_body(doc, APPROVED["intro"])
    add_table(doc, ["기간", "경험", "역할·배운 점"], [
        ["2022년", "창원 PCB 제조업체 생산 현장", "낯을 가리고 질문하지 못했던 고비를 겪음. 먼저 인사하고 제때 질문하는 행동으로 바꿈."],
        ["2026년", "IT·정보보안 교육과정", "문제를 하나하나 확인하고, 동료의 말을 듣고 해결 방법을 설명하는 학습 방식을 이어 감."],
        ["2026.08–09", "리추얼·과제 기록", "기록을 연 날과 실천했다고 남긴 날을 통해 반복 행동을 확인함."],
    ], [2.4, 4.2, 10.2])
    add_heading(doc, "관심 분야", 2)
    add_bullet(doc, "웹 애플리케이션: 로그인과 인증 흐름을 이해하고 사용자를 보호하는 기능 만들기")
    add_bullet(doc, "정보보안: 문제를 기록하고 확인하며, 필요한 사람과 함께 해결하는 일")
    add_heading(doc, "대표작", 2)
    for work in APPROVED["works"]:
        add_bullet(doc, f"{work['number']}번 {work['type']} — {work['title']}: {work['description']}")

    add_heading(doc, "2. 자기소개서", 1)
    for chapter in APPROVED["story"]:
        add_body(doc, chapter["text"])
    add_heading(doc, "세 능력이 드러난 장면", 2)
    add_table(doc, ["능력", "장면", "현재의 행동"], [
        [strength["name"], strength["before"], strength["now"]]
        for strength in APPROVED["strengths"]
    ], [2.8, 6.1, 7.9])

    add_heading(doc, "3. 경력기술서", 1)
    add_body(doc, "경력기술서는 확인된 경험을 상황·행동·결과로 정리했습니다. 다른 사람의 이름은 기록에 남기지 않았습니다.")
    add_table(doc, ["과제·기간", "능력", "상황 · 행동 · 결과"], [
        ["창원 PCB 제조업체 / 2022년", "자기조절력", "상황: 낯을 가리고 선배들이 어렵게 느껴져 모르는 것도 쉽게 질문하지 못함. 행동: 먼저 밝게 인사하고 모르는 것은 그때그때 질문함. 결과: 혼자 고민하는 대신 필요한 내용을 제때 묻는 순서를 익힘."],
        ["IT·정보보안 교육과정 / 2026년", "자기동기력", "상황: 실습 중 문제가 생김. 행동: 문제를 대충 넘기지 않고 하나씩 정독하며 해결 방법을 찾음. 결과: 어려움을 성장의 계기로 받아들이고 학습을 이어 가는 방식을 만듦."],
        ["리추얼·과제 기록 / 2026.08–09", "대인관계력", "상황: 과제를 진행하며 동료와 함께 막힌 부분을 풀어야 했음. 행동: 먼저 해결한 부분은 설명하고 다른 사람의 의견과 도움은 끝까지 들음. 결과: 경청·도움·설명의 반복을 기록으로 남김."],
        ["10번 논문 / 완료", "자기동기력", "상황: 로그인 공격의 분산 방식과 속도에 따른 탐지 정책을 비교해야 했음. 행동: 다섯 가지 모의 시나리오와 세 가지 탐지 정책을 비교함. 결과: 공격 탐지와 정상 사용자 보호의 균형을 살핌."],
    ], [3.2, 2.3, 11.3])
    add_heading(doc, "지원 방향", 2)
    add_body(doc, "아직 완성된 전문가라고 말할 수는 없습니다. 대신 모르는 것을 숨기지 않고, 기록을 확인하고, 필요한 사람에게 질문하고, 해결한 방법을 다시 설명하는 방식으로 배우고 있습니다. 웹 애플리케이션과 정보보안 분야에서 이 태도를 더 정확한 지식과 책임감 있는 협업으로 발전시키고 싶습니다.")

    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    doc.core_properties.title = "김태빈 지원 서류 묶음"
    doc.core_properties.subject = "이력서·자기소개서·경력기술서"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
